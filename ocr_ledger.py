from __future__ import annotations

import base64
import hashlib
import io
import json
import math
import os
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import requests
import xlrd
import xlwt
from PIL import Image, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from xlutils.copy import copy as copy_workbook
from xlwt.Cell import BlankCell, NumberCell, StrCell


JOB_URL = os.environ.get("PADDLEOCR_JOB_URL", "https://paddleocr.aistudio-app.com/api/v2/ocr/jobs")
OCR_MODEL = os.environ.get("PADDLEOCR_MODEL", "PP-OCRv6")
VISION_LLM_API_URL = os.environ.get("VISION_LLM_API_URL", "").strip()
VISION_LLM_API_KEY = os.environ.get("VISION_LLM_API_KEY", "").strip()
VISION_LLM_MODEL = os.environ.get("VISION_LLM_MODEL", "").strip()
TEMPLATE_PATH = Path(os.environ.get("LEDGER_TEMPLATE_PATH") or Path(__file__).resolve().parent / "templates" / "default_template.xls")
ZERO_WIDTH = "\u200b"


@dataclass
class OcrItem:
    text: str
    score: float
    box: list[float]
    x: float
    y: float


@dataclass
class ParsedRow:
    source_file: str
    source_row: int
    source_image: str = ""
    crop_image: str = ""
    name: str = ""
    id_no: str = ""
    address: str = ""
    phone: str = ""
    amount_text: str = ""
    amount: float | None = None
    quantity: float | None = None
    anomalies: list[str] = field(default_factory=list)
    raw_text: str = ""
    ocr_note: str = ""


def clean_text(text: Any) -> str:
    value = str(text or "").replace("×", "X").replace("脳", "X")
    value = value.replace("（", "(").replace("）", ")")
    return re.sub(r"\s+", "", value)


def normalize_cell_text(value: Any) -> str:
    return clean_text(value).strip("，,。.()（）-:：")


def normalize_id(value: Any) -> str:
    text = clean_text(value).upper().replace("Ｏ", "0").replace("O", "0")
    match = re.search(r"\d{17}[\dX]", text)
    return match.group(0) if match else text


def normalize_phone(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def normalize_amount_text(value: Any) -> str:
    return clean_text(value).replace("＋", "+").strip(".。")


def parse_amount(text: str, anomalies: list[str]) -> float | None:
    clean = normalize_amount_text(text).replace(",", "").replace("，", "").replace("元", "")
    if not clean:
        return None
    if re.fullmatch(r"\d+(?:\.\d+)?", clean):
        return float(clean)
    if re.fullmatch(r"\d+(?:\.\d+)?(?:\+\d+(?:\.\d+)?)+", clean):
        anomalies.append("amount expression summed")
        return sum(float(part) for part in clean.split("+"))
    match = re.search(r"\d+(?:\.\d+)?", clean)
    if match:
        anomalies.append("amount contained noise; numeric part extracted")
        return float(match.group(0))
    anomalies.append("amount could not be parsed")
    return None


def validate_row(row: ParsedRow) -> None:
    if not row.name or len(row.name) > 8:
        row.anomalies.append("name needs review")
    if not re.fullmatch(r"\d{17}[\dX]", row.id_no or ""):
        row.anomalies.append("ID number needs review")
    if not row.address:
        row.anomalies.append("address needs review")
    if not re.fullmatch(r"\d{11}", row.phone or ""):
        row.anomalies.append("phone needs review")
    if row.amount is None:
        row.anomalies.append("amount needs review")


def submit_ocr_job(image_path: Path, token: str) -> str:
    headers = {"Authorization": f"bearer {token}"}
    payload = {"useDocOrientationClassify": False, "useDocUnwarping": False, "useTextlineOrientation": False}
    data = {"model": OCR_MODEL, "optionalPayload": json.dumps(payload, ensure_ascii=False)}
    with image_path.open("rb") as file_handle:
        response = requests.post(JOB_URL, headers=headers, data=data, files={"file": file_handle}, timeout=60)
    if response.status_code != 200:
        raise RuntimeError(f"OCR submit failed: HTTP {response.status_code} {response.text[:300]}")
    return response.json()["data"]["jobId"]


def wait_for_result(job_id: str, token: str, poll_seconds: int = 5) -> str:
    headers = {"Authorization": f"bearer {token}"}
    while True:
        response = requests.get(f"{JOB_URL}/{job_id}", headers=headers, timeout=60)
        if response.status_code != 200:
            raise RuntimeError(f"OCR polling failed: HTTP {response.status_code} {response.text[:300]}")
        data = response.json()["data"]
        if data["state"] == "done":
            return data["resultUrl"]["jsonUrl"]
        if data["state"] == "failed":
            raise RuntimeError(f"OCR job failed: {data.get('errorMsg', 'unknown error')}")
        time.sleep(poll_seconds)


def download_jsonl(json_url: str) -> str:
    response = requests.get(json_url, timeout=120)
    response.raise_for_status()
    response.encoding = "utf-8"
    return response.text


def parse_jsonl(jsonl_text: str) -> dict[str, Any]:
    for line in jsonl_text.splitlines():
        if line.strip():
            return json.loads(line)
    raise ValueError("empty OCR result")


def average_confidence(parsed: dict[str, Any]) -> float:
    scores = parsed["result"]["ocrResults"][0]["prunedResult"].get("rec_scores", [])
    return sum(float(score) for score in scores) / len(scores) if scores else 0.0


def fingerprint_result(parsed: dict[str, Any]) -> str:
    pruned = parsed["result"]["ocrResults"][0]["prunedResult"]
    rows = [[clean_text(text), [round(float(v), 1) for v in box]] for text, box in zip(pruned.get("rec_texts", []), pruned.get("rec_boxes", []))]
    return hashlib.sha256(json.dumps(rows, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def run_ocr_three_times(image_path: Path, token: str, output_dir: Path, log: list[str]) -> tuple[str, dict[str, Any]]:
    runs: list[dict[str, Any]] = []
    for run_index in range(1, 4):
        log.append(f"{image_path.name}: OCR run {run_index}/3")
        job_id = submit_ocr_job(image_path, token)
        json_url = wait_for_result(job_id, token)
        jsonl_text = download_jsonl(json_url)
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / f"{image_path.stem}_run{run_index}.jsonl").write_text(jsonl_text, encoding="utf-8")
        parsed = parse_jsonl(jsonl_text)
        runs.append({"run_index": run_index, "jsonl_text": jsonl_text, "fingerprint": fingerprint_result(parsed), "confidence": average_confidence(parsed)})
    grouped: dict[str, list[dict[str, Any]]] = {}
    for run in runs:
        grouped.setdefault(run["fingerprint"], []).append(run)
    best_group = max(grouped.values(), key=lambda group: (len(group), max(item["confidence"] for item in group)))
    chosen = max(best_group, key=lambda item: item["confidence"])
    note = f"three OCR runs {'matched' if len(grouped) == 1 else 'differed'}; selected run {chosen['run_index']}"
    log.append(f"{image_path.name}: {note}")
    return chosen["jsonl_text"], {"chosen_run": chosen["run_index"], "compare_note": note}


def items_from_parsed(parsed: dict[str, Any]) -> list[OcrItem]:
    pruned = parsed["result"]["ocrResults"][0]["prunedResult"]
    items: list[OcrItem] = []
    for text, score, box in zip(pruned.get("rec_texts", []), pruned.get("rec_scores", []), pruned.get("rec_boxes", [])):
        box = list(map(float, box))
        items.append(OcrItem(clean_text(text), float(score), box, (box[0] + box[2]) / 2, (box[1] + box[3]) / 2))
    return items


def cluster_rows(items: list[OcrItem]) -> dict[int, list[OcrItem]]:
    if not items:
        return {}
    width = max(item.box[2] for item in items)
    anchors = {int(item.text): item.y for item in items if re.fullmatch(r"\d{1,2}", item.text) and 1 <= int(item.text) <= 30 and item.x < width * 0.16}
    if not anchors:
        y_values = sorted(item.y for item in items if re.search(r"\d{8,}", item.text) or (item.x > width * 0.62 and re.search(r"\d{3,}", item.text)))
        anchors = {index + 1: value for index, value in enumerate(y_values[:30])}
    rows = {key: [] for key in anchors}
    if not anchors:
        return rows
    centers = sorted(anchors.items())
    gaps = [centers[i + 1][1] - centers[i][1] for i in range(len(centers) - 1) if centers[i + 1][1] > centers[i][1]]
    tolerance = max(28.0, (sorted(gaps)[len(gaps) // 2] if gaps else 52.0) * 0.5)
    for item in items:
        row_no, row_y = min(anchors.items(), key=lambda pair: abs(pair[1] - item.y))
        if abs(row_y - item.y) <= tolerance:
            rows.setdefault(row_no, []).append(item)
    for row_items in rows.values():
        row_items.sort(key=lambda item: (item.x, item.y))
    return rows


def parse_one_row(row_items: list[OcrItem], source_file: str, source_row: int, unit_price: float) -> ParsedRow:
    row = ParsedRow(source_file=source_file, source_row=source_row)
    if not row_items:
        return row
    width = max(item.box[2] for item in row_items)
    row.raw_text = " | ".join(item.text for item in row_items)
    main_text = "".join(item.text for item in row_items if width * 0.10 <= item.x <= width * 0.62 and not re.fullmatch(r"\d{1,2}", item.text))
    amount_items = [item for item in row_items if item.x >= width * 0.62 and re.search(r"\d", item.text)]
    id_match = re.search(r"\d{17}[\dXx]|\d{18}", main_text)
    if id_match:
        row.name = normalize_cell_text(main_text[: id_match.start()])
        row.id_no = normalize_id(id_match.group(0))
        rest = main_text[id_match.end() :]
    else:
        row.name = normalize_cell_text(main_text)
        rest = ""
    phone_matches = list(re.finditer(r"\d{8,13}", rest))
    if phone_matches:
        match = phone_matches[-1]
        row.address = normalize_cell_text(rest[: match.start()])
        row.phone = normalize_phone(match.group(0))[:11]
    else:
        row.address = normalize_cell_text(rest)
    if amount_items:
        amount_items.sort(key=lambda item: (-item.score, item.x))
        row.amount_text = normalize_amount_text(amount_items[0].text)
    row.amount = parse_amount(row.amount_text, row.anomalies)
    row.quantity = round(row.amount / unit_price, 2) if row.amount is not None and unit_price > 0 else None
    validate_row(row)
    if any(item.score < 0.75 for item in row_items):
        row.anomalies.append("low-confidence OCR in row")
    return row


def parse_rows_from_jsonl(jsonl_text: str, source_file: str, unit_price: float) -> list[ParsedRow]:
    rows = []
    for source_row, row_items in sorted(cluster_rows(items_from_parsed(parse_jsonl(jsonl_text))).items()):
        row = parse_one_row(row_items, source_file, source_row, unit_price)
        if row.name or row.id_no or row.address or row.phone or row.amount_text:
            rows.append(row)
    return rows


def vision_llm_enabled() -> bool:
    return bool(VISION_LLM_API_URL and VISION_LLM_API_KEY and VISION_LLM_MODEL)


def make_row_crop(image_path: Path, row_items: list[OcrItem]) -> Image.Image:
    with Image.open(image_path) as image:
        image = ImageOps.exif_transpose(image).convert("RGB")
        width, height = image.size
        if row_items:
            top = max(0, int(min(item.y for item in row_items) - 42))
            bottom = min(height, int(max(item.y for item in row_items) + 42))
        else:
            top, bottom = 0, height
        crop = image.crop((0, top, int(width * 0.9), bottom))
        if crop.size[1] < 140:
            scale = 140 / max(1, crop.size[1])
            crop = crop.resize((int(crop.size[0] * scale), 140), Image.Resampling.LANCZOS)
        return crop


def image_to_data_url(image: Image.Image, quality: int = 78) -> str:
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG", quality=quality, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")


def rows_from_llm_payload(payload: dict[str, Any], base_rows: list[ParsedRow], source_file: str, unit_price: float) -> list[ParsedRow]:
    raw_rows = payload.get("rows", []) if isinstance(payload.get("rows", []), list) else [payload]
    by_row = {row.source_row: row for row in base_rows}
    result: list[ParsedRow] = []
    for item in raw_rows:
        if not isinstance(item, dict):
            continue
        source_row = int(item.get("source_row") or 0)
        base = by_row.get(source_row)
        row = ParsedRow(source_file=source_file, source_row=source_row or (base.source_row if base else 0))
        if base:
            row.raw_text = base.raw_text
            row.crop_image = base.crop_image
            row.source_image = base.source_image
        row.name = normalize_cell_text(item.get("name", base.name if base else ""))
        row.id_no = normalize_id(item.get("id_no", base.id_no if base else ""))
        row.address = normalize_cell_text(item.get("address", base.address if base else ""))
        row.phone = normalize_phone(item.get("phone", base.phone if base else ""))[:11]
        row.amount_text = normalize_amount_text(item.get("amount_text", base.amount_text if base else ""))
        row.amount = parse_amount(row.amount_text, row.anomalies)
        row.quantity = round(row.amount / unit_price, 2) if row.amount is not None and unit_price > 0 else None
        note = normalize_cell_text(item.get("review_note", ""))
        validate_row(row)
        if note:
            row.anomalies.append(f"LLM_VL review: {note}")
        result.append(row)
    return sorted(result, key=lambda row: row.source_row)


def refine_rows_with_llm(jsonl_text: str, rows: list[ParsedRow], image_path: Path, source_file: str, unit_price: float, log: list[str], crop_dir: Path | None = None) -> list[ParsedRow]:
    parsed_rows = cluster_rows(items_from_parsed(parse_jsonl(jsonl_text)))
    crop_dir = crop_dir or image_path.parent / "row_crops"
    crop_dir.mkdir(parents=True, exist_ok=True)
    for row in rows:
        row.source_image = str(image_path)
        crop = make_row_crop(image_path, parsed_rows.get(row.source_row, []))
        crop_path = crop_dir / f"{image_path.stem}_row{row.source_row:02d}.jpg"
        crop.save(crop_path, format="JPEG", quality=86, optimize=True)
        row.crop_image = str(crop_path)
    if not vision_llm_enabled():
        log.append(f"{source_file}: LLM_VL not configured; skipped visual review")
        return rows
    refined: list[ParsedRow] = []
    headers = {"Authorization": f"Bearer {VISION_LLM_API_KEY}", "Content-Type": "application/json"}
    for row in rows:
        crop = make_row_crop(image_path, parsed_rows.get(row.source_row, []))
        user_payload = {
            "task": "Review one cropped row from a structured handwritten ledger and return JSON for name, id_no, address, phone, amount_text.",
            "source_row": row.source_row,
            "ocr_row": row.raw_text,
            "preliminary_row": row.__dict__,
        }
        payload = {
            "model": VISION_LLM_MODEL,
            "messages": [
                {"role": "system", "content": "You are a strict multimodal OCR reviewer. Use the image as primary evidence and OCR text as auxiliary evidence. Return JSON only."},
                {"role": "user", "content": [{"type": "text", "text": json.dumps(user_payload, ensure_ascii=False)}, {"type": "image_url", "image_url": {"url": image_to_data_url(crop)}}]},
            ],
            "temperature": 0.0,
            "response_format": {"type": "json_object"},
        }
        try:
            response = requests.post(VISION_LLM_API_URL, headers=headers, json=payload, timeout=90)
            if response.status_code != 200:
                raise RuntimeError(f"HTTP {response.status_code}: {response.text[:160]}")
            content = response.json()["choices"][0]["message"]["content"]
            text = content[content.find("{") : content.rfind("}") + 1]
            refined_row = rows_from_llm_payload({"rows": [json.loads(text)]}, [row], source_file, unit_price)[0]
            refined_row.source_image = row.source_image
            refined_row.crop_image = row.crop_image
            refined.append(refined_row)
        except Exception as exc:
            row.anomalies.append(f"LLM_VL failed; kept OCR result: {exc}")
            refined.append(row)
    log.append(f"{source_file}: LLM_VL reviewed {len(refined)} rows")
    return refined


def as_text(value: str) -> str:
    return f"{ZERO_WIDTH}{value}" if value else ""


def typed_number_or_text(value: Any) -> Any:
    text = str(value or "").strip()
    if re.fullmatch(r"\d+", text):
        return int(text)
    if re.fullmatch(r"\d+\.\d+", text):
        return float(text)
    return text


def copied_row(sheet: xlwt.Worksheet, row_index: int) -> xlwt.Row.Row:
    return sheet.row(row_index)


def template_xf_index(read_sheet: xlrd.sheet.Sheet, write_sheet: xlwt.Worksheet, row_index: int, col_index: int) -> int:
    rows = getattr(write_sheet, "_Worksheet__rows")
    if row_index in rows:
        cells = getattr(rows[row_index], "_Row__cells")
        if col_index in cells:
            return int(cells[col_index].xf_idx)
    for source_row in range(min(row_index, read_sheet.nrows - 1), -1, -1):
        if source_row in rows:
            cells = getattr(rows[source_row], "_Row__cells")
            if col_index in cells:
                return int(cells[col_index].xf_idx)
    return 0


def write_preserving_template_style(write_book: xlwt.Workbook, write_sheet: xlwt.Worksheet, read_sheet: xlrd.sheet.Sheet, row_index: int, col_index: int, value: Any) -> None:
    row = copied_row(write_sheet, row_index)
    xf_index = template_xf_index(read_sheet, write_sheet, row_index, col_index)
    if value is None or value == "":
        row.insert_cell(col_index, BlankCell(row_index, col_index, xf_index))
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        row.insert_cell(col_index, NumberCell(row_index, col_index, xf_index, float(value)))
    else:
        row.insert_cell(col_index, StrCell(row_index, col_index, xf_index, write_book.add_str(str(value))))


def write_template_workbook(rows: list[ParsedRow], template_path: Path, output_path: Path, unit_price: float, product_type: str, defaults: dict[str, str], purchase_date: str = "") -> None:
    read_book = xlrd.open_workbook(str(template_path), formatting_info=True)
    write_book = copy_workbook(read_book)
    sheet = write_book.get_sheet(0)
    read_sheet = read_book.sheet_by_index(0)
    start_row = 3
    max_rows = max(read_sheet.nrows - start_row, len(rows), 165)
    for row_index in range(start_row, start_row + max_rows):
        for col in range(18):
            write_preserving_template_style(write_book, sheet, read_sheet, row_index, col, "")
    for index, row in enumerate(rows, start=1):
        excel_row = start_row + index - 1
        values = {0: index, 1: purchase_date, 2: row.name, 3: as_text(row.id_no), 4: row.address, 5: as_text(row.phone), 6: product_type, 7: typed_number_or_text(defaults.get("grade", "1")), 8: defaults.get("unit", "公斤"), 9: row.quantity if row.quantity is not None else "", 10: unit_price, 11: row.amount if row.amount is not None else row.amount_text, 12: defaults.get("origin", "本地"), 13: defaults.get("self_sold", "是"), 17: ""}
        for col, value in values.items():
            write_preserving_template_style(write_book, sheet, read_sheet, excel_row, col, value)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    write_book.save(str(output_path))


def make_xlsx_styles() -> dict[str, Any]:
    side = Side(style="thin", color="D9E2F3")
    return {"border": Border(left=side, right=side, top=side, bottom=side), "center": Alignment(horizontal="center", vertical="center", wrap_text=True), "wrap": Alignment(horizontal="left", vertical="center", wrap_text=True)}


def write_xlsx_workbook(rows: list[ParsedRow], output_path: Path, unit_price: float, product_type: str, defaults: dict[str, str], purchase_date: str = "") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Extracted"
    headers = ["序号", "收购日期", "出售人姓名", "出售人身份证号码", "出售人详细地址", "出售人联系电话", "农产品名称", "等级/规格", "计量单位", "收购数量", "收购单价", "收购总金额", "产品产地", "是否自产", "收购人签字", "入库单号", "付款方式", "备注"]
    for column_index, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column_index, value=header)
    styles = make_xlsx_styles()
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.border = styles["border"]
        cell.alignment = styles["center"]
    for index, row in enumerate(rows, start=1):
        values = [index, purchase_date, row.name, row.id_no, row.address, row.phone, product_type, typed_number_or_text(defaults.get("grade", "1")), defaults.get("unit", "公斤"), row.quantity, unit_price, row.amount if row.amount is not None else row.amount_text, defaults.get("origin", "本地"), defaults.get("self_sold", "是"), "", "", "", ""]
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=index + 1, column=column_index, value=value)
            cell.border = styles["border"]
            cell.alignment = styles["wrap"]
            if column_index in {4, 6} and value:
                cell.value = str(value)
                cell.number_format = "@"
            if column_index in {10, 11, 12} and isinstance(value, (int, float)):
                cell.number_format = "0.00"
    widths = [8, 12, 16, 24, 28, 18, 12, 10, 10, 14, 14, 16, 12, 12, 12, 12, 12, 32]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def review_payload_to_rows(payload: dict[str, Any], unit_price: float) -> list[ParsedRow]:
    rows: list[ParsedRow] = []
    for index, item in enumerate(payload.get("rows", []), start=1):
        row = ParsedRow(source_file=str(item.get("source_file", "")), source_row=int(item.get("source_row", index) or index), source_image=str(item.get("source_image", "")), crop_image=str(item.get("crop_image", "")), name=normalize_cell_text(item.get("name", "")), id_no=normalize_id(item.get("id_no", "")), address=normalize_cell_text(item.get("address", "")), phone=normalize_phone(item.get("phone", "")), amount_text=normalize_amount_text(item.get("amount_text", item.get("amount", ""))), anomalies=list(item.get("anomalies") or []), raw_text=str(item.get("raw_text", "")), ocr_note=str(item.get("ocr_note", "")))
        row.amount = parse_amount(row.amount_text, [])
        row.quantity = round(row.amount / unit_price, 2) if row.amount is not None and unit_price > 0 else None
        rows.append(row)
    return rows


def refresh_outputs_from_review_payload(payload: dict[str, Any]) -> dict[str, Any]:
    settings = dict(payload.get("settings") or {})
    unit_price = float(settings.get("unit_price") or 2.42)
    product_type = str(settings.get("product_type") or "小麦")
    defaults = dict(settings.get("defaults") or {})
    defaults.setdefault("origin", "本地")
    defaults.setdefault("self_sold", "是")
    defaults.setdefault("grade", "1")
    defaults.setdefault("unit", "公斤")
    purchase_date = str(settings.get("purchase_date") or "")
    template_path = Path(settings.get("template_path") or TEMPLATE_PATH)
    output_path = Path(payload.get("output") or Path.cwd() / f"ledger_{product_type}_{time.strftime('%Y%m%d_%H%M%S')}.xls")
    review_workbook_path = Path(payload.get("review_workbook") or output_path.with_name(f"ledger_{product_type}_review_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"))
    rows = review_payload_to_rows(payload, unit_price)
    write_template_workbook(rows, template_path, output_path, unit_price, product_type, defaults, purchase_date)
    write_xlsx_workbook(rows, review_workbook_path, unit_price, product_type, defaults, purchase_date)
    payload.update({"output": str(output_path), "review_workbook": str(review_workbook_path), "row_count": len(rows), "flagged_count": sum(1 for row in rows if row.anomalies), "rows": [row.__dict__ for row in rows], "saved_at": time.strftime("%Y-%m-%d %H:%M:%S"), "settings": {"template_path": str(template_path), "unit_price": unit_price, "product_type": product_type, "defaults": defaults, "purchase_date": purchase_date}})
    return payload


def process_images(image_paths: list[Path], token: str, template_path: Path, output_dir: Path, unit_price: float, product_type: str, defaults: dict[str, str], purchase_date: str = "", log: list[str] | None = None) -> dict[str, Any]:
    log = log or []
    all_rows: list[ParsedRow] = []
    comparisons: list[dict[str, Any]] = []
    raw_dir = output_dir / "ocr_raw"
    crop_dir = output_dir / "row_crops"
    for image_path in image_paths:
        jsonl_text, comparison = run_ocr_three_times(image_path, token, raw_dir, log)
        rows = parse_rows_from_jsonl(jsonl_text, image_path.name, unit_price)
        rows = refine_rows_with_llm(jsonl_text, rows, image_path, image_path.name, unit_price, log, crop_dir)
        for row in rows:
            row.ocr_note = comparison["compare_note"]
        all_rows.extend(rows)
        comparisons.append(comparison)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"ledger_{product_type}_{stamp}.xls"
    review_workbook_path = output_dir / f"ledger_{product_type}_review_{stamp}.xlsx"
    write_template_workbook(all_rows, template_path, output_path, unit_price, product_type, defaults, purchase_date)
    write_xlsx_workbook(all_rows, review_workbook_path, unit_price, product_type, defaults, purchase_date)
    payload = {"output": str(output_path), "review_workbook": str(review_workbook_path), "row_count": len(all_rows), "flagged_count": sum(1 for row in all_rows if row.anomalies), "settings": {"template_path": str(template_path), "unit_price": unit_price, "product_type": product_type, "defaults": defaults, "purchase_date": purchase_date}, "comparisons": comparisons, "rows": [row.__dict__ for row in all_rows], "log": log}
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "审核明细.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def safe_price(value: str) -> float:
    try:
        price = float(value)
    except ValueError as exc:
        raise ValueError("unit price must be numeric") from exc
    if not math.isfinite(price) or price <= 0:
        raise ValueError("unit price must be greater than zero")
    return price
